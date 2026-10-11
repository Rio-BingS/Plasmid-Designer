"""整理包命名与写出内容的安全性（网页端整理包与离线脚本共用 sequencing_report）

覆盖：超长名字截断保留扩展名、不同质粒安全化后撞名时加 ~N 区分（短名不变）。
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.sequencing_report import (  # noqa: E402
    BatchArchive,
    _safe_component,
    _unique_components,
    batch_archive_entries,
)


# ---------------------------------------------------------------- L2 命名


def test_safe_component_short_names_unchanged():
    # 未超长的名字与此前实现逐字一致（网页整理包里的既有名字不变）
    assert _safe_component("T1.ab1") == "T1.ab1"
    assert _safe_component("../../ESCAPE") == "ESCAPE"
    assert _safe_component("..") == "unnamed"
    assert _safe_component("a/b") == "a_b"
    assert _safe_component("x" * 80) == "x" * 80


def test_safe_component_truncation_keeps_extension():
    name = "L" * 120 + "1.fasta"
    out = _safe_component(name)
    assert len(out) <= 80
    assert out.endswith(".fasta")
    assert out.startswith("L" * 70)
    # 截断不留 Windows 不接受的结尾空格/点
    assert _safe_component("A" * 74 + "  " + "B" * 30 + ".ab1") == "A" * 74 + ".ab1"


def test_safe_component_long_name_without_extension():
    out = _safe_component("N" * 200)
    assert out == "N" * 80
    # 「扩展名」过长（不像扩展名）时按原样截断
    out = _safe_component("a." + "b" * 100)
    assert len(out) == 80 and out.startswith("a.")


def test_unique_components_disambiguates_deterministically():
    names = ["a/b", "a:b", "a_b"]
    m = _unique_components(names)
    # 安全化前后不变的 'a_b' 保留原名，其余按原始名排序加 ~2、~3
    assert m["a_b"] == "a_b"
    assert sorted([m["a/b"], m["a:b"]]) == ["a_b~2", "a_b~3"]
    assert len(set(m.values())) == 3
    # 与输入顺序无关
    assert _unique_components(list(reversed(names))) == m


def test_unique_components_long_names_stay_within_limit():
    a, b = "L" * 100 + "A", "L" * 100 + "B"
    m = _unique_components([a, b])
    assert m[a] == "L" * 80
    assert m[b] == "L" * 78 + "~2"
    assert all(len(v) <= 80 for v in m.values())


def test_unique_components_case_insensitive_and_short_untouched():
    m = _unique_components(["P1", "P2", "x/y", "X_Y"])
    assert m["P1"] == "P1" and m["P2"] == "P2"
    assert {m["x/y"], m["X_Y"]} == {"X_Y", "x_y~2"}


def _group(plasmid, ref_name=None, read_names=()):
    ref = ({"file": {"name": ref_name, "bytes": b">r\nACGT\n"}, "how": "质粒名"}
           if ref_name else None)
    reads = [{"file": {"name": n, "bytes": n.encode()}, "how": "引物列"}
             for n in read_names]
    return {"plasmid": plasmid, "clone": None, "reference": ref, "reads": reads}


def _item(plasmid):
    return {"plasmid": plasmid, "clone": None, "conclusion": "无法判定：测试",
            "analysis_id": None, "status": "skipped", "read_count": 1,
            "coverage_percent": None, "variant_count": 0, "reference_name": None}


def _entries(groups):
    items = [_item(g["plasmid"]) for g in groups]
    payload = {"items": items, "unmatched": [], "clone_mode": False}
    return [a for a, _d in batch_archive_entries(payload, groups, [], None, None,
                                                 None, {})]


def test_archive_separates_plasmids_that_sanitize_to_same_folder():
    groups = [_group("a/b", read_names=["r1.ab1"]),
              _group("a:b", read_names=["r2.ab1"])]
    arcs = _entries(groups)
    assert "测序整理/a_b/无法判定/r1.ab1" in arcs
    assert "测序整理/a_b~2/无法判定/r2.ab1" in arcs
    # 两份报告各在自己的文件夹，不再出现 测序分析报告 (2).md
    assert "测序整理/a_b/无法判定/测序分析报告.md" in arcs
    assert "测序整理/a_b~2/无法判定/测序分析报告.md" in arcs
    assert not any("(2)" in a for a in arcs)


def test_archive_long_plasmid_names_keep_reference_extension():
    p1, p2 = "L" * 100 + "1", "L" * 100 + "2"
    groups = [_group(p1, ref_name=p1 + ".fasta", read_names=["a.ab1"]),
              _group(p2, ref_name=p2 + ".fasta", read_names=["b.ab1"])]
    arcs = _entries(groups)
    refs = [a for a in arcs if a.endswith(".fasta")]
    assert len(refs) == 2
    folders = {a.split("/")[1] for a in refs}
    assert folders == {"L" * 80, "L" * 78 + "~2"}
    assert all(len(seg) <= 80 for a in arcs for seg in a.split("/"))


def test_archive_short_names_unchanged():
    groups = [_group("TST", ref_name="TST.fasta", read_names=["T1.ab1"])]
    arcs = _entries(groups)
    assert "测序整理/TST/TST.fasta" in arcs
    assert "测序整理/TST/无法判定/T1.ab1" in arcs
    assert BatchArchive(["TST"]).folder_of(groups[0]) == "TST"


# ---------------------------------------------------------------- L3 CSV / Markdown 注入

import csv  # noqa: E402
import importlib.util  # noqa: E402
import io  # noqa: E402
from pathlib import Path  # noqa: E402

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from app.sequencing_report import (  # noqa: E402
    _group_report_md,
    csv_guard,
    csv_unguard,
)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from abif_utils import make_ab1  # noqa: E402

_REPO = Path(__file__).resolve().parents[1]
_spec = importlib.util.spec_from_file_location(
    "batch_sequencing_report_l3", _REPO / "scripts" / "batch_sequencing_report.py")
batch = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(batch)


@pytest.mark.parametrize("raw", [
    "=HYPERLINK(1)|p.ab1", "+1", "-X/正确/a.ab1", "@SUM(A1)", "\tx", "\rx",
    "'=x", "''+y", "'plain", "plain", "", "a=b",
])
def test_csv_guard_neutralizes_and_round_trips(raw):
    g = csv_guard(raw)
    assert csv_unguard(g) == raw
    if raw[:1] in ("=", "+", "-", "@", "\t", "\r"):
        assert g == "'" + raw
    if raw in ("plain", "", "a=b", "'plain"):
        assert g == raw          # 安全内容不加前缀
    assert g[:1] not in ("=", "+", "-", "@", "\t", "\r")
    assert csv_guard(12.5) == 12.5


def _csv_rows(text: str) -> list:
    return list(csv.DictReader(io.StringIO(text.lstrip("\ufeff"))))


def test_web_manifest_cells_neutralized_and_report_table_escaped():
    evil = "=HYPERLINK(1)|p.ab1"
    groups = [_group("-P", read_names=[evil, "ok.ab1"])]
    items = [_item("-P")]
    payload = {"items": items, "unmatched": [], "clone_mode": False}
    entries = dict(batch_archive_entries(payload, groups, [], None, None, None, {}))
    rows = _csv_rows(entries["测序整理/整理清单.csv"])
    evil_row = next(r for r in rows if "HYPERLINK" in r["文件"])
    assert evil_row["文件"] == "'" + evil
    assert evil_row["质粒"] == "'-P"
    assert not evil_row["归档路径"].startswith("'")   # 有根前缀，不需防护
    ok_row = next(r for r in rows if r["文件"] == "ok.ab1")
    assert ok_row["结论"] == "无法判定：测试"
    for r in rows:
        for v in r.values():
            assert v[:1] not in ("=", "+", "-", "@", "\t", "\r"), r
    report = next(v for k, v in entries.items() if k.endswith("测序分析报告.md"))
    line = next(ln for ln in report.splitlines() if "HYPERLINK" in ln)
    # 表格仍是 3 列：未转义的 | 只有列分隔符那 4 个
    assert line.replace("\\|", "").count("|") == 4
    assert "=HYPERLINK(1)\\|p.ab1" in line


def test_report_table_newline_in_name_stays_one_row():
    item = {"plasmid": "P", "conclusion": "无法判定：测试", "reference_name": None}
    md = _group_report_md(item, None, [{"name": "bad\nname|x.ab1", "kind": "测序 read",
                                        "how": "引物列"}])
    rows = [ln for ln in md.splitlines() if ln.startswith("| bad")]
    assert rows == ["| bad name\\|x.ab1 | 测序 read | 引物列 |"]


def test_report_plain_names_unchanged():
    item = {"plasmid": "P", "conclusion": "无法判定：测试", "reference_name": None}
    md = _group_report_md(item, None, [{"name": "T1.ab1", "kind": "测序 read",
                                        "how": "引物列"}])
    assert "| T1.ab1 | 测序 read | 引物列 |" in md.splitlines()


REF = ("ATGAAACGT" * 12 + "TAA")


def _offline_inputs(data_dir: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒名称", "测序引物", "测序结果"])
    ws.append(["-EVIL", "@T1", None])
    wb.save(data_dir / "测序.xlsx")
    (data_dir / "-EVIL.fasta").write_text(f">E\n{REF}\n", encoding="utf-8")
    (data_dir / "@T1.ab1").write_bytes(make_ab1(REF, [40] * len(REF)))


def _run(monkeypatch, data_dir):
    monkeypatch.setattr(sys, "argv", [
        "batch_sequencing_report.py", "--data-dir", str(data_dir), "--no-write-back"])
    batch.main()


def test_offline_manifest_neutralized_and_rerun_still_tracks_paths(tmp_path, monkeypatch):
    data_dir = tmp_path / "results"
    data_dir.mkdir()
    _offline_inputs(data_dir)
    _run(monkeypatch, data_dir)
    out = data_dir / "测序分析"
    rows = _csv_rows((out / "整理清单.csv").read_text(encoding="utf-8-sig"))
    by_file = {csv_unguard(r["文件"]): r for r in rows}
    assert set(by_file) == {"-EVIL.fasta", "@T1.ab1"}
    read = by_file["@T1.ab1"]
    assert read["文件"] == "'@T1.ab1" and read["质粒"] == "'-EVIL"
    assert read["归档路径"].startswith("'-EVIL/")
    arc = csv_unguard(read["归档路径"])
    assert (out / arc).is_file()

    # 用户改动输出目录里的副本 → 重跑另存 (2)，清单仍记实际路径（带防护前缀）
    (out / arc).write_bytes(b"user edited")
    _run(monkeypatch, data_dir)
    rows = _csv_rows((out / "整理清单.csv").read_text(encoding="utf-8-sig"))
    read = next(r for r in rows if csv_unguard(r["文件"]) == "@T1.ab1")
    arc2 = csv_unguard(read["归档路径"])
    assert read["归档路径"].startswith("'-EVIL/") and arc2.endswith("@T1 (2).ab1")
    import hashlib
    assert hashlib.md5((out / arc2).read_bytes()).hexdigest() == read["MD5"]


def test_legacy_manifest_fallback_reads_guarded_paths(tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    text = ("\ufeff质粒,克隆,文件,类型,结论,匹配方式,归档路径,大小KB,MD5\r\n"
            "'-EVIL,,'@T1.ab1,测序 read,,引物列,'-EVIL/正确/@T1.ab1,0.1,abc\r\n"
            "P,,a.ab1,测序 read,,引物列,P/正确/a.ab1,0.1,def\r\n")
    (out / "整理清单.csv").write_text(text, encoding="utf-8", newline="")
    assert batch.load_previous_outputs(out) == {
        "-EVIL/正确/@T1.ab1": "abc", "P/正确/a.ab1": "def"}


# ---------------------------------------------------------------- L4 Excel 回填重跑

from app.sequencing_report import merge_result_cell  # noqa: E402

TOOL_LINES = [
    "合格：与设计一致",
    "合格：编码区蛋白一致；编码区外 1 处确证差异（ori）",
    "不合格：GFP 蛋白与设计不一致（移码）；测序覆盖 87%",
    "疑似混合：T1.ab1 检出 5 处双峰，无法自动判定，建议重新挑单克隆复测",
    "疑似混合样品：T1.ab1 检出 5 处双峰",              # 旧版本措辞
    "无法判定：部分比对（T1.ab1），未比对部分未验证",
    "分析失败",
    "分析失败：参考序列过短",
    "已整理 1 个测序文件，但缺少同名参考图谱（.dna/.gb/.fasta），无法自动比对",
    "未在文件夹中找到该质粒的测序文件或图谱",
    "未找到克隆 C1 的图谱或测序文件",
    "已找到参考图谱，但未匹配到它的测序文件（.ab1）",
    "已找到参考图谱，但未匹配到克隆 C1 的测序文件（.ab1）",
]


@pytest.mark.parametrize("prev", TOOL_LINES)
def test_merge_replaces_previous_tool_conclusion(prev):
    assert merge_result_cell(prev, "合格：与设计一致") == "合格：与设计一致"


def test_merge_replaces_stacked_legacy_lines_keeps_user_text():
    old = ("客户要求加急\n"
           "已整理 1 个测序文件，但缺少同名参考图谱（.dna/.gb/.fasta），无法自动比对\n"
           "合格：与设计一致；测序覆盖 87%")
    new = merge_result_cell(old, "不合格：GFP 蛋白与设计不一致")
    assert new == "客户要求加急\n不合格：GFP 蛋白与设计不一致"
    # 再跑一次结果稳定（幂等）
    assert merge_result_cell(new, "不合格：GFP 蛋白与设计不一致") == new


def test_merge_user_text_only_and_empty():
    assert merge_result_cell(None, "合格：与设计一致") == "合格：与设计一致"
    assert merge_result_cell("  ", "合格：与设计一致") == "合格：与设计一致"
    assert merge_result_cell("  待复测  ", "合格：与设计一致") == "待复测\n合格：与设计一致"
    # 用户手写的「合格」（无冒号）、提及结论的句子不当作工具输出
    assert merge_result_cell("合格", "无法判定：x") == "合格\n无法判定：x"
    assert merge_result_cell("老板说合格：先发货", "无法判定：x") == "老板说合格：先发货\n无法判定：x"


def test_merge_flattens_multiline_conclusion():
    out = merge_result_cell("", "分析失败：line1\nline2")
    assert out == "分析失败：line1 line2"
    assert merge_result_cell(out, "合格：与设计一致") == "合格：与设计一致"


def test_offline_rerun_replaces_backfilled_result(tmp_path, monkeypatch):
    data_dir = tmp_path / "results"
    data_dir.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒名称", "测序引物", "测序结果"])
    ws.append(["TST", "T1", "加急"])
    wb.save(data_dir / "测序.xlsx")
    (data_dir / "T1.ab1").write_bytes(make_ab1(REF, [40] * len(REF)))

    def run():
        monkeypatch.setattr(sys, "argv", [
            "batch_sequencing_report.py", "--data-dir", str(data_dir)])
        batch.main()
        return openpyxl.load_workbook(data_dir / "测序.xlsx").worksheets[0].cell(
            row=2, column=3).value

    first = run()          # 缺图谱 → 无法自动比对
    assert first.startswith("加急\n已整理 1 个测序文件，但缺少同名参考图谱")
    (data_dir / "TST.fasta").write_text(f">TST\n{REF}\n", encoding="utf-8")
    assert run() == "加急\n合格：与设计一致"
    assert run() == "加急\n合格：与设计一致"     # 不再层层叠加
