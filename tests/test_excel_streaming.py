"""信息表流式解析回归锁（审查 02 M5：Excel 用 read_only 流式解析）

此前 load_excel 一律以完整模式 load_workbook：1.3 MB 的 xlsx（5000 行 ×
130 列）解析 3.5 s、峰值 310 MB，行列上限要等整表建完对象才检查。现在
解析走 openpyxl read_only 逐行读取、边读边查上限；只有结论回填才另行以
可写模式打开。解析结果须与完整模式逐字段一致，且不信任文件自带的
dimension 声明。
"""

import io
import re
import zipfile
from datetime import datetime

import openpyxl
import pytest

from core.sanger import batch as sb


def _full_mode_reference(xlsx: bytes):
    """改造前 load_excel 的完整模式解析（逐字段对照用）"""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb.worksheets[0]
    header_row, cols = None, {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        for c in row:
            if c.value and "质粒" in str(c.value):
                header_row = c.row
                break
        if header_row:
            break
    for c in ws[header_row]:
        v = str(c.value or "")
        if "质粒" in v:
            cols["plasmid"] = c.column
        elif "引物" in v:
            cols["primer"] = c.column
        elif "结果" in v:
            cols["result"] = c.column
        elif "克隆" in v or "样品" in v:
            cols["clone"] = c.column
    cols.setdefault("primer", cols["plasmid"] + 1)
    cols.setdefault("result", cols["primer"] + 1)
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = sb._cell_str(ws.cell(row=r, column=cols["plasmid"]).value)
        prim = sb._cell_str(ws.cell(row=r, column=cols["primer"]).value)
        if not name and not prim:
            continue
        rows.append({
            "row": r, "name": name,
            "clone": sb._cell_str(ws.cell(row=r, column=cols["clone"]).value)
            if "clone" in cols else "",
            "primers": [p.strip() for p in sb.PRIMER_SPLIT.split(prim) if p.strip()],
            "existing_result": ws.cell(row=r, column=cols["result"]).value,
        })
    return header_row, cols, rows


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_variants():
    # 1) 标题行在表头之上、克隆号列、数值单元格、空行与稀疏行
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["某测序公司 送样单"])
    ws.append([])
    ws.append(["序号", "克隆号", "质粒名称", "测序引物", "测序结果", "备注"])
    ws.append([1, 123.0, "MX", "M13F；M13R", None, "x"])
    ws.append([])
    ws.append([2, 124, "MX", "M13F, T7", "已有结论", None])
    ws.cell(row=12, column=3, value="pUC19")              # 稀疏：中间缺行
    ws.cell(row=12, column=4, value="T7.ab1\nSP6.ab1")
    ws.cell(row=13, column=5, value=datetime(2026, 1, 2))  # 只有结果列
    ws.cell(row=14, column=4, value=5)                    # 只有引物（数值）
    yield _save(wb)
    # 2) 无引物/结果列（缺省取相邻列）、表头在第 1 行
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒"])
    ws.append(["P1", "a.ab1;b.ab1", "旧"])
    ws.append(["P2"])
    yield _save(wb)
    # 3) 合并单元格 + 公式
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒名称", "测序引物", "测序结果"])
    ws.append(["MX", "A", "=1+1"])
    ws.append(["", "B", None])
    ws.merge_cells("A2:A3")
    yield _save(wb)


@pytest.mark.parametrize("xlsx", list(_sheet_variants()))
def test_streaming_parse_matches_full_mode(xlsx):
    assert sb.parse_excel(xlsx) == _full_mode_reference(xlsx)
    wb, header_row, cols, rows = sb.load_excel(xlsx)
    assert (header_row, cols, rows) == _full_mode_reference(xlsx)
    assert not getattr(wb, "read_only", False)            # 回填用的仍是可写工作簿


def test_parse_uses_read_only_and_load_excel_reopens_writable(monkeypatch):
    modes = []
    real = openpyxl.load_workbook

    def spy(*a, **kw):
        modes.append(kw.get("read_only", False))
        return real(*a, **kw)

    monkeypatch.setattr(openpyxl, "load_workbook", spy)
    xlsx = next(_sheet_variants())
    sb.parse_excel(xlsx)
    assert modes == [True]
    modes.clear()
    sb.load_excel(xlsx)
    assert modes == [True, False]


def _rewrite_dimension(xlsx: bytes, ref: str) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as src, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                if "<dimension" in text:
                    text = re.sub(r'<dimension ref="[^"]*"\s*/>', f'<dimension ref="{ref}"/>', text)
                else:
                    text = text.replace("<sheetData", f'<dimension ref="{ref}"/><sheetData', 1)
                data = text.encode("utf-8")
            dst.writestr(info, data)
    return out.getvalue()


def _plain_sheet(n):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒名称", "测序引物", "测序结果"])
    for i in range(n):
        ws.append([f"P{i}", f"T{i}", None])
    return _save(wb)


def test_understated_dimension_does_not_drop_rows():
    lying = _rewrite_dimension(_plain_sheet(30), "A1:C2")
    _h, _c, rows = sb.parse_excel(lying)
    assert [r["name"] for r in rows] == [f"P{i}" for i in range(30)]


def test_row_cap_counts_real_rows_not_declared_dimension(monkeypatch):
    monkeypatch.setattr(sb, "MAX_EXCEL_ROWS", 50)
    # 声明只有 2 行、实际 100 行：仍按实际行数拒绝
    with pytest.raises(ValueError, match="行列数"):
        sb.parse_excel(_rewrite_dimension(_plain_sheet(100), "A1:C2"))
    # 声明 10 万行、实际 10 行：按实际内容放行
    _h, _c, rows = sb.parse_excel(_rewrite_dimension(_plain_sheet(10), "A1:C100000"))
    assert len(rows) == 10


def test_column_cap(monkeypatch):
    monkeypatch.setattr(sb, "MAX_EXCEL_COLS", 5)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["质粒名称", "测序引物", "测序结果"])
    ws.cell(row=2, column=9, value="x")
    with pytest.raises(ValueError, match="行列数"):
        sb.parse_excel(_save(wb))


def test_missing_header_still_reported():
    wb = openpyxl.Workbook()
    wb.active.append(["名称", "引物"])
    with pytest.raises(ValueError, match="表头"):
        sb.parse_excel(_save(wb))
